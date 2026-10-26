"""FastAPI 应用:REST API + 静态前端托管。"""
from __future__ import annotations

import io
import traceback

import pandas as pd

from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

from . import __version__
from .codegen import code_for
from .core.datasets import STORE, Dataset, _safe
from .core.registry import REGISTRY
from .core.results import ModelError, Result
from .models import CATEGORIES  # noqa: F401  触发注册
from .processing import OPERATIONS
from .samples import SAMPLES, create as create_sample

app = FastAPI(title="OpenSPSS", version=__version__)

import os  # noqa: E402
import sys  # noqa: E402

WEB_DIR = os.path.join(os.path.dirname(__file__), "web")
if not os.path.isdir(WEB_DIR) and getattr(sys, "_MEIPASS", None):
    # PyInstaller 单文件模式:静态资源在解包目录
    cand = os.path.join(sys._MEIPASS, "open_spss", "web")  # noqa: SLF001
    if os.path.isdir(cand):
        WEB_DIR = cand


@app.exception_handler(ModelError)
async def model_error_handler(_req, exc: ModelError):
    return JSONResponse(status_code=400, detail=str(exc))


class RunPayload(BaseModel):
    model: str
    dataset: str
    slots: dict = Field(default_factory=dict)
    options: dict = Field(default_factory=dict)


class ProcessPayload(BaseModel):
    op: str
    params: dict = Field(default_factory=dict)


class ExportPayload(BaseModel):
    title: str = "分析结果"
    tables: list[dict]


class ColumnPayload(BaseModel):
    column: str
    dtype: str


class RenamePayload(BaseModel):
    old: str
    new: str


class EditPayload(BaseModel):
    """数据工作台编辑:cell / add_row / del_row / add_col / del_col"""
    action: str
    row: int | None = None
    column: str | None = None
    name: str | None = None
    value: str | None = None
    dtype: str = "cat"


class LabelsPayload(BaseModel):
    column: str
    label: str = ""
    vals: dict = Field(default_factory=dict)


class WeightPayload(BaseModel):
    column: str | None = None


class MergePayload(BaseModel):
    other_id: str
    how: str = "append"          # append 纵向追加 / join 横向关联
    on: str = ""
    join: str = "inner"          # join 模式:inner / left


# ---------------------------------------------------------------- 元信息
@app.get("/api/health")
def health():
    import sys
    return {"status": "ok", "version": __version__, "models": len(REGISTRY.all()),
            "frozen": bool(getattr(sys, "frozen", False))}


@app.post("/api/shutdown")
def shutdown():
    """打包版在前端页脚提供「退出」;开发模式一般不会调用。"""
    import os
    import threading

    threading.Timer(0.4, lambda: os._exit(0)).start()
    return {"ok": True}


@app.get("/api/catalog")
def catalog():
    return {"categories": CATEGORIES, "total": len(REGISTRY.all())}


# ---------------------------------------------------------------- 数据集
@app.get("/api/datasets")
def list_datasets():
    return {"datasets": STORE.list(), "samples": [{"key": k, "label": v[0]} for k, v in SAMPLES.items()]}


@app.post("/api/upload")
async def upload(file: UploadFile = File(...)):
    content = await file.read()
    if len(content) > 60 * 1024 * 1024:
        raise ModelError("文件超过 60MB 限制")
    df, meta = STORE.read_upload(file.filename or "data.csv", content)
    ds = STORE.add(file.filename.rsplit(".", 1)[0], df, source="upload", meta=meta)
    return ds.summary()


@app.post("/api/datasets/sample/{key}")
def load_sample(key: str):
    if key not in SAMPLES:
        raise ModelError(f"未知示例数据集: {key}")
    ds = create_sample(key)
    return ds.summary()


@app.get("/api/datasets/{ds_id}")
def dataset_detail(ds_id: str, page: int = 1, page_size: int = 50):
    ds = STORE.get(ds_id)
    d = ds.summary()
    d["preview"] = ds.preview(page, page_size)
    return d


@app.delete("/api/datasets/{ds_id}")
def delete_dataset(ds_id: str):
    STORE.delete(ds_id)
    return {"ok": True}


@app.put("/api/datasets/{ds_id}/columns")
def set_column(ds_id: str, payload: ColumnPayload):
    ds = STORE.get(ds_id)
    ds.set_dtype(payload.column, payload.dtype)
    return {"ok": True, "columns": ds.col_meta()}


@app.post("/api/datasets/{ds_id}/rename")
def rename_column(ds_id: str, payload: RenamePayload):
    ds = STORE.get(ds_id)
    ds.rename(payload.old, payload.new)
    return {"ok": True, "columns": ds.col_meta()}


@app.post("/api/datasets/{ds_id}/process")
def process(ds_id: str, payload: ProcessPayload):
    ds = STORE.get(ds_id)
    fn = OPERATIONS.get(payload.op)
    if fn is None:
        raise ModelError(f"未知的数据处理操作: {payload.op}")
    try:
        msg = fn(ds, **payload.params)
    except TypeError as e:
        raise ModelError(f"参数错误: {e}") from e
    ds.note(msg)  # note 内部落盘:数据处理对列/行的改动一并持久化
    d = ds.summary()
    d["message"] = msg
    d["preview"] = ds.preview(1, 50)
    return d


# ---------------------------------------------------------------- 数据工作台
@app.post("/api/datasets/{ds_id}/edit")
def edit_dataset(ds_id: str, payload: EditPayload):
    ds = STORE.get(ds_id)
    actions = {
        "cell": lambda: ds.set_cell(payload.row, payload.column, payload.value),
        "add_row": ds.add_row,
        "del_row": lambda: ds.del_row(payload.row),
        "add_col": lambda: ds.add_col(payload.name, payload.dtype),
        "del_col": lambda: ds.del_col(payload.column),
    }
    fn = actions.get(payload.action)
    if fn is None:
        raise ModelError(f"未知编辑操作: {payload.action}")
    fn()
    ds.note({"cell": f"修改单元格 {payload.column}[{payload.row}]",
             "add_row": "新增 1 行", "del_row": f"删除第 {payload.row + 1} 行",
             "add_col": f"新增列「{payload.name}」", "del_col": f"删除列「{payload.column}」"}[payload.action])
    d = ds.summary()
    d["preview"] = ds.preview(1, 50)
    return d


@app.post("/api/datasets/{ds_id}/undo")
def undo_dataset(ds_id: str):
    ds = STORE.get(ds_id)
    if not ds.undo():
        raise ModelError("没有可撤销的操作")
    ds.note("撤销上一步数据修改")
    d = ds.summary()
    d["preview"] = ds.preview(1, 50)
    return d


@app.post("/api/datasets/{ds_id}/redo")
def redo_dataset(ds_id: str):
    ds = STORE.get(ds_id)
    if not ds.redo():
        raise ModelError("没有可重做的操作")
    ds.note("重做数据修改")
    d = ds.summary()
    d["preview"] = ds.preview(1, 50)
    return d


@app.post("/api/datasets/{ds_id}/labels")
def set_labels(ds_id: str, payload: LabelsPayload):
    ds = STORE.get(ds_id)
    ds.set_labels(payload.column, payload.label, payload.vals or None)
    return {"ok": True, "columns": ds.col_meta()}


@app.post("/api/datasets/{ds_id}/weight")
def set_weight(ds_id: str, payload: WeightPayload):
    ds = STORE.get(ds_id)
    ds.set_weight(payload.column or None)
    return {"ok": True, "weightCol": ds.weight_col}


@app.post("/api/datasets/{ds_id}/merge")
def merge_dataset(ds_id: str, payload: MergePayload):
    ds = STORE.get(ds_id)
    other = STORE.get(payload.other_id)
    if payload.how == "append":
        merged = pd.concat([ds.df, other.df], ignore_index=True)
        name = f"{ds.name}+{other.name}"
    elif payload.how == "join":
        if not payload.on:
            raise ModelError("请选择关联键列")
        if payload.on not in ds.df.columns or payload.on not in other.df.columns:
            raise ModelError(f"关联键「{payload.on}」需同时存在于两个数据集")
        if payload.join not in ("inner", "left"):
            raise ModelError("关联方式仅支持 inner / left")
        merged = pd.merge(ds.df, other.df, on=payload.on, how=payload.join,
                          suffixes=("", "_2"))
        name = f"{ds.name}⋈{other.name}"
    else:
        raise ModelError("合并方式仅支持 append / join")
    new_ds = STORE.add(name, merged, source="merge")
    new_ds.note(f"由「{ds.name}」与「{other.name}」"
                f"{'纵向追加' if payload.how == 'append' else '按「' + payload.on + '」' + ('左连' if payload.join == 'left' else '内联')}合并")
    return new_ds.summary()


# ---------------------------------------------------------------- 运行模型
@app.post("/api/run")
def run_model(payload: RunPayload):
    try:
        spec = REGISTRY.get(payload.model)
    except KeyError:
        raise HTTPException(404, f"模型不存在: {payload.model}")
    ds = STORE.get(payload.dataset)
    try:
        slots, opts = REGISTRY.validate(payload.model, payload.slots, payload.options)
        # 类型校验
        dtype_map = {c["name"]: c["dtype"] for c in ds.columns}
        for s in spec.slots:
            val = slots.get(s.key)
            cols = val if isinstance(val, list) else ([val] if val else [])
            for c in cols:
                if c not in dtype_map:
                    raise ModelError(f"变量「{c}」不在当前数据集中")
                if not s.accepts(dtype_map[c]):
                    tname = {"quant": "定量", "cat": "定类", "ordinal": "定级"}.get(dtype_map[c], dtype_map[c])
                    raise ModelError(f"「{s.label}」不接收{tname}变量「{c}」,请在数据页修改变量类型或换选变量")
        result = spec.runner(ds, slots, opts)
    except ModelError as e:
        raise HTTPException(400, str(e))
    except RecursionError:
        raise HTTPException(400, "数据规模过大导致计算中断,请抽样后重试")
    except Exception as e:  # noqa: BLE001
        traceback.print_exc()
        raise HTTPException(500, f"计算出错: {type(e).__name__}: {e}")
    out = result.to_dict()
    out["model"] = {"id": spec.id, "name": spec.name}
    code = code_for(spec.id, ds, slots, opts)
    if code:
        out["code"] = code
    return out


# ---------------------------------------------------------------- 导出
def _tables_to_xlsx(payload: ExportPayload) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="EEF2FF")
    head_font = Font(bold=True, color="312E81", size=11)
    title_font = Font(bold=True, size=13, color="1E1B4B")
    thin = Side(style="thin", color="D9DEE8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for i, t in enumerate(payload.tables):
        ws = wb.create_sheet(f"表{i + 1}"[:31])
        ncol = max(len(t.get("headers") or []), max((len(r) for r in t.get("rows", [])), default=0), 1)
        ws.cell(1, 1, t.get("title") or f"表{i + 1}").font = title_font
        for j, h in enumerate(t.get("headers") or [], 1):
            c = ws.cell(2, j, str(h))
            c.fill, c.font, c.border = head_fill, head_font, border
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for r, row in enumerate(t.get("rows") or [], 3):
            for j, v in enumerate(row[:ncol], 1):
                c = ws.cell(r, j, _safe(v) if not isinstance(v, (int, float, str)) or isinstance(v, bool) else v)
                c.border = border
                c.alignment = Alignment(horizontal="center", vertical="center")
        if t.get("note"):
            nc = ws.cell(ws.max_row + 2, 1, "注:" + t["note"])
            nc.font = Font(italic=True, size=10, color="64748B")
        for j in range(1, ncol + 1):
            width = max(10, min(38, max(len(str(ws.cell(r, j).value or "")) for r in range(1, ws.max_row + 1)) + 4))
            ws.column_dimensions[get_column_letter(j)].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@app.post("/api/export/xlsx")
def export_xlsx(payload: ExportPayload):
    data = _tables_to_xlsx(payload)
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=result.xlsx"},
    )


# ---------------------------------------------------------------- 前端
@app.get("/favicon.ico")
def favicon():
    return FileResponse(os.path.join(WEB_DIR, "favicon.svg"), media_type="image/svg+xml")


app.mount("/", StaticFiles(directory=WEB_DIR, html=True), name="web")
