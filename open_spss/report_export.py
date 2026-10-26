"""综合分析报告导出:Word(python-docx)与 PDF(reportlab CID 中文字体)。

由前端收集报告内容(标题/解读/表格/图表 PNG),后端只做排版,
不依赖数据库或会话状态,同一份载荷可同时导出两种格式。
"""
from __future__ import annotations

import base64
import io

from pydantic import BaseModel, Field


class ReportTable(BaseModel):
    title: str = ""
    headers: list[str] = Field(default_factory=list)
    rows: list[list] = Field(default_factory=list)
    note: str = ""


class ReportChart(BaseModel):
    title: str = ""
    png: str = ""            # dataURL 或裸 base64


class ReportItem(BaseModel):
    title: str
    meta: str = ""
    slots: str = ""
    paragraphs: list[str] = Field(default_factory=list)
    tables: list[ReportTable] = Field(default_factory=list)
    charts: list[ReportChart] = Field(default_factory=list)


class ReportPayload(BaseModel):
    title: str = "数据分析报告"
    generated: str = ""
    dark: bool = False
    items: list[ReportItem] = Field(default_factory=list)


def _png_bytes(dataurl: str) -> bytes | None:
    if not dataurl:
        return None
    b64 = dataurl.split(",", 1)[1] if dataurl.startswith("data:") else dataurl
    try:
        return base64.b64decode(b64)
    except Exception:  # noqa: BLE001
        return None


# ---------------------------------------------------------------- Word
def build_docx(p: ReportPayload) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    h = doc.add_heading(p.title, level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if p.generated:
        meta = doc.add_paragraph(p.generated)
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta.runs[0].font.size = Pt(9)
        meta.runs[0].font.color.rgb = RGBColor(0x8B, 0x94, 0xAB)
    doc.add_paragraph(f"共 {len(p.items)} 项分析 · 由 OpenSPSS 生成")

    for it in p.items:
        doc.add_heading(it.title, level=1)
        if it.meta:
            mp = doc.add_paragraph(it.meta)
            mp.runs[0].font.size = Pt(9)
            mp.runs[0].font.color.rgb = RGBColor(0x8B, 0x94, 0xAB)
        if it.slots:
            sp = doc.add_paragraph("变量配置:" + it.slots)
            sp.runs[0].font.size = Pt(9)
        for para in it.paragraphs:
            doc.add_paragraph(_plain(para))
        for t in it.tables:
            if t.title:
                doc.add_heading(t.title, level=2)
            if not t.headers:
                continue
            tb = doc.add_table(rows=1, cols=len(t.headers))
            tb.style = "Table Grid"
            for j, htxt in enumerate(t.headers):
                cell = tb.rows[0].cells[j]
                cell.text = str(htxt)
                for run in cell.paragraphs[0].runs:
                    run.font.bold = True
                    run.font.size = Pt(9)
            for row in t.rows:
                cells = tb.add_row().cells
                for j, v in enumerate(row[:len(t.headers)]):
                    cells[j].text = str(v if v is not None else "")
                    for run in cells[j].paragraphs[0].runs:
                        run.font.size = Pt(9)
            if t.note:
                np_ = doc.add_paragraph("注:" + t.note)
                np_.runs[0].font.size = Pt(8.5)
                np_.runs[0].font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        for c in it.charts:
            img = _png_bytes(c.png)
            if img:
                if c.title:
                    doc.add_heading(c.title, level=2)
                doc.add_picture(io.BytesIO(img), width=Cm(15))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _plain(md: str) -> str:
    """解读文本里的 **bold** 与 `code` 标记在 Word 中退化为纯文本。"""
    import re
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", md)
    return text.replace("`", "")


# ---------------------------------------------------------------- PDF
def build_pdf(p: ReportPayload) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import (Image, PageBreak, Paragraph, SimpleDocTemplate,
                                    Spacer, Table)

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    ink = colors.HexColor("#1B2338") if not p.dark else colors.HexColor("#DEE4F5")
    muted = colors.HexColor("#64748B") if not p.dark else colors.HexColor("#8B96B5")
    line = colors.HexColor("#D6DBE8") if not p.dark else colors.HexColor("#303A5E")
    head_bg = colors.HexColor("#F1F5F9") if not p.dark else colors.HexColor("#1C2340")

    st_title = ParagraphStyle("t", fontName="STSong-Light", fontSize=20, alignment=1,
                              textColor=ink, spaceAfter=6)
    st_meta = ParagraphStyle("m", fontName="STSong-Light", fontSize=9, alignment=1,
                             textColor=muted, spaceAfter=14)
    st_h1 = ParagraphStyle("h1", fontName="STSong-Light", fontSize=14, textColor=ink,
                           spaceBefore=14, spaceAfter=6, borderPadding=3)
    st_h2 = ParagraphStyle("h2", fontName="STSong-Light", fontSize=11, textColor=ink,
                           spaceBefore=8, spaceAfter=4)
    st_body = ParagraphStyle("b", fontName="STSong-Light", fontSize=10, leading=16,
                             textColor=ink, spaceAfter=4)
    st_small = ParagraphStyle("s", fontName="STSong-Light", fontSize=8.5, leading=12,
                              textColor=muted)
    st_cell = ParagraphStyle("c", fontName="STSong-Light", fontSize=8, leading=11,
                             textColor=ink)

    story: list = [Paragraph(p.title, st_title)]
    if p.generated:
        story.append(Paragraph(p.generated, st_meta))
    for idx, it in enumerate(p.items):
        if idx:
            story.append(PageBreak())
        story.append(Paragraph(it.title, st_h1))
        if it.meta:
            story.append(Paragraph(it.meta, st_small))
        if it.slots:
            story.append(Paragraph("变量配置:" + it.slots, st_small))
        for para in it.paragraphs:
            story.append(Paragraph(_plain(para), st_body))
        for t in it.tables:
            if t.title:
                story.append(Paragraph(t.title, st_h2))
            if not t.headers:
                continue
            data = [[Paragraph(str(h), st_cell) for h in t.headers]]
            for row in t.rows:
                data.append([Paragraph(str(v if v is not None else ""), st_cell)
                             for v in row[:len(t.headers)]])
            width = 17 * cm / len(t.headers)
            tb = Table(data, colWidths=[width] * len(t.headers), repeatRows=1)
            tb.setStyle([("GRID", (0, 0), (-1, -1), 0.4, line),
                         ("BACKGROUND", (0, 0), (-1, 0), head_bg),
                         ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                         ("TOPPADDING", (0, 0), (-1, -1), 2),
                         ("BOTTOMPADDING", (0, 0), (-1, -1), 2)])
            story.append(tb)
            if t.note:
                story.append(Paragraph("注:" + t.note, st_small))
        for c in it.charts:
            img = _png_bytes(c.png)
            if img:
                if c.title:
                    story.append(Paragraph(c.title, st_h2))
                story.append(Spacer(1, 4))
                story.append(Image(io.BytesIO(img), width=15 * cm, height=15 * cm * 0.55,
                                   kind="proportional"))
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, title=p.title).build(story)
    return buf.getvalue()
